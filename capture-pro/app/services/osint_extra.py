"""OSINT Extra — Profile, Metadata, API integrations."""
import os
import io
import re
import json
import base64
import socket
from datetime import datetime

try:
    import requests
    HAS_REQUESTS = True
except ImportError:
    HAS_REQUESTS = False

# Import basic OSINT
try:
    from app.services.osint_pro import (
        ip_intel, email_breach, domain_whois, dns_lookup,
        ssl_info, http_headers, phone_lookup, username_check
    )
    HAS_BASIC = True
except ImportError:
    HAS_BASIC = False


def _key(service):
    """Ambil API key dari env."""
    return os.environ.get(service.upper() + "_API_KEY", "")


# ============================================================
# 1. TARGET PROFILE AGGREGATOR
# ============================================================
def target_profile(target):
    """Aggregate semua info + risk score."""
    if not HAS_BASIC:
        return {"target": target, "error": "Basic OSINT service tidak tersedia"}

    result = {"target": target, "lookup_at": datetime.now().isoformat()}

    try:
        if re.match(r"^\d{1,3}(\.\d{1,3}){3}$", target):
            result["type"] = "ip"
            result["ip_intel"] = ip_intel(target)
        elif "@" in target:
            result["type"] = "email"
            result["breach"] = email_breach(target)
            if "@" in target:
                domain = target.split("@")[1]
                result["domain_whois"] = domain_whois(domain)
        elif re.match(r"^[+\d\s\-()]+$", target) and len(target) >= 8:
            result["type"] = "phone"
            result["phone"] = phone_lookup(target)
        elif "." in target and " " not in target:
            result["type"] = "domain"
            result["whois"] = domain_whois(target)
            result["dns"] = dns_lookup(target)
            result["ssl"] = ssl_info(target)
        else:
            result["type"] = "username"
            result["username"] = username_check(target)

        # Risk score
        result["risk_score"] = _calc_risk(result)
    except Exception as e:
        result["error"] = str(e)

    return result


def _calc_risk(p):
    """Hitung risk score 0-100."""
    score = 0
    # Email breach
    breach = p.get("breach", {})
    if breach.get("status") == "BREACHED":
        score += min(50, breach.get("count", 0) * 10)
    # VPN
    proxy = p.get("ip_intel", {}).get("proxycheck", {})
    if proxy.get("proxy") == "yes":
        score += 30
    # Missing SSL
    ssl = p.get("ssl", {})
    if ssl.get("error"):
        score += 20
    return min(100, score)


# ============================================================
# 2. DOCUMENT METADATA EXTRACTOR
# ============================================================
def extract_metadata(file_bytes, filename):
    """Extract metadata dari PDF/DOCX/XLSX."""
    result = {"filename": filename, "size_bytes": len(file_bytes),
              "lookup_at": datetime.now().isoformat()}
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            meta = {}
            if reader.metadata:
                for k, v in dict(reader.metadata).items():
                    meta[str(k)] = str(v)
            result["metadata"] = meta
            result["pages"] = len(reader.pages)
        except Exception as e:
            result["error"] = f"PDF: {e}"

    elif ext in ("docx", "doc"):
        try:
            import docx
            d = docx.Document(io.BytesIO(file_bytes))
            p = d.core_properties
            result["metadata"] = {
                "author": p.author or "",
                "last_modified_by": p.last_modified_by or "",
                "created": str(p.created) if p.created else "",
                "modified": str(p.modified) if p.modified else "",
                "title": p.title or "",
                "subject": p.subject or "",
                "keywords": p.keywords or "",
                "category": p.category or "",
                "comments": p.comments or "",
                "revision": str(p.revision) if p.revision else "",
                "content_status": p.content_status or "",
            }
            result["paragraphs"] = len(d.paragraphs)
        except Exception as e:
            result["error"] = f"DOCX: {e}"

    elif ext in ("xlsx", "xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
            p = wb.properties
            result["metadata"] = {
                "creator": p.creator or "",
                "last_modified_by": p.lastModifiedBy or "",
                "created": str(p.created) if p.created else "",
                "modified": str(p.modified) if p.modified else "",
                "title": p.title or "",
                "subject": p.subject or "",
                "description": p.description or "",
                "keywords": p.keywords or "",
                "category": p.category or "",
            }
            result["sheets"] = wb.sheetnames
        except Exception as e:
            result["error"] = f"XLSX: {e}"

    else:
        result["error"] = f"Format .{ext} tidak didukung (pakai pdf/docx/xlsx)"

    return result


# ============================================================
# 3. SHODAN
# ============================================================
def shodan_lookup(ip):
    """Cari info device dari Shodan."""
    result = {"ip": ip, "service": "shodan", "lookup_at": datetime.now().isoformat()}
    key = _key("shodan")
    if not key:
        result["error"] = "SHODAN_API_KEY belum diset di .env"
        result["register_url"] = "https://account.shodan.io/register"
        return result
    try:
        r = requests.get(f"https://api.shodan.io/shodan/host/{ip}?key={key}", timeout=20)
        d = r.json()
        if "error" in d:
            result["error"] = d["error"]
        else:
            result["data"] = {
                "ip": d.get("ip_str"),
                "org": d.get("org"),
                "isp": d.get("isp"),
                "country": d.get("country_name"),
                "city": d.get("city"),
                "os": d.get("os"),
                "ports": d.get("ports", []),
                "hostnames": d.get("hostnames", []),
                "domains": d.get("domains", []),
                "vulns": list(d.get("vulns", {}).keys()),
                "last_update": d.get("last_update"),
            }
    except Exception as e:
        result["error"] = str(e)
    return result


# ============================================================
# 4. VIRUSTOTAL
# ============================================================
def virustotal_lookup(target, type_="ip"):
    """Cek reputasi di VirusTotal."""
    result = {"target": target, "type": type_, "service": "virustotal",
              "lookup_at": datetime.now().isoformat()}
    key = _key("virustotal")
    if not key:
        result["error"] = "VIRUSTOTAL_API_KEY belum diset di .env"
        result["register_url"] = "https://www.virustotal.com/gui/join-us"
        return result
    try:
        if type_ == "ip":
            url = f"https://www.virustotal.com/api/v3/ip_addresses/{target}"
        elif type_ == "domain":
            url = f"https://www.virustotal.com/api/v3/domains/{target}"
        elif type_ == "url":
            url_id = base64.urlsafe_b64encode(target.encode()).decode().rstrip("=")
            url = f"https://www.virustotal.com/api/v3/urls/{url_id}"
        else:
            result["error"] = f"Tipe {type_} tidak didukung"
            return result

        r = requests.get(url, headers={"x-apikey": key}, timeout=20)
        d = r.json()
        if "error" in d:
            result["error"] = d["error"].get("message", str(d["error"]))
        else:
            attrs = d.get("data", {}).get("attributes", {})
            stats = attrs.get("last_analysis_stats", {})
            result["data"] = {
                "reputation": attrs.get("reputation", 0),
                "malicious": stats.get("malicious", 0),
                "suspicious": stats.get("suspicious", 0),
                "harmless": stats.get("harmless", 0),
                "undetected": stats.get("undetected", 0),
                "total_engines": sum(stats.values()) if stats else 0,
                "country": attrs.get("country"),
                "as_owner": attrs.get("as_owner"),
            }
    except Exception as e:
        result["error"] = str(e)
    return result


# ============================================================
# 5. URLSCAN
# ============================================================
def urlscan_submit(url):
    """Submit URL ke URLScan.io."""
    result = {"url": url, "service": "urlscan", "lookup_at": datetime.now().isoformat()}
    key = _key("urlscan")
    if not key:
        result["error"] = "URLSCAN_API_KEY belum diset di .env"
        result["register_url"] = "https://urlscan.io/user/signup"
        return result
    try:
        r = requests.post("https://urlscan.io/api/v1/scan/",
                          headers={"API-Key": key, "Content-Type": "application/json"},
                          json={"url": url, "visibility": "unlisted"}, timeout=30)
        d = r.json()
        if "uuid" in d:
            result["uuid"] = d["uuid"]
            result["web_url"] = d.get("result", "")
            result["message"] = "Scan berjalan. Cek hasil 10-30 detik di web_url"
        else:
            result["error"] = d.get("message", str(d))
    except Exception as e:
        result["error"] = str(e)
    return result


# ============================================================
# 6. HUNTER.IO
# ============================================================
def hunter_search(domain):
    """Cari email dari domain via Hunter.io."""
    result = {"domain": domain, "service": "hunter", "lookup_at": datetime.now().isoformat()}
    key = _key("hunter")
    if not key:
        result["error"] = "HUNTER_API_KEY belum diset di .env"
        result["register_url"] = "https://hunter.io/users/sign_up"
        return result
    try:
        r = requests.get(f"https://api.hunter.io/v2/domain-search?domain={domain}&api_key={key}",
                         timeout=20)
        d = r.json().get("data", {})
        result["data"] = {
            "organization": d.get("organization"),
            "pattern": d.get("pattern"),
            "total": len(d.get("emails", [])),
            "emails": [
                {"email": e.get("value"), "type": e.get("type"),
                 "first_name": e.get("first_name"), "last_name": e.get("last_name"),
                 "position": e.get("position")}
                for e in d.get("emails", [])[:30]
            ],
        }
    except Exception as e:
        result["error"] = str(e)
    return result


# ============================================================
# 7. ABUSEIPDB
# ============================================================
def abuseipdb_check(ip):
    """Cek abuse report untuk IP."""
    result = {"ip": ip, "service": "abuseipdb", "lookup_at": datetime.now().isoformat()}
    key = _key("abuseipdb")
    if not key:
        result["error"] = "ABUSEIPDB_API_KEY belum diset di .env"
        result["register_url"] = "https://www.abuseipdb.com/register"
        return result
    try:
        r = requests.get("https://api.abuseipdb.com/api/v2/check",
                         params={"ipAddress": ip, "maxAgeInDays": 90},
                         headers={"Key": key, "Accept": "application/json"},
                         timeout=15)
        d = r.json().get("data", {})
        result["data"] = {
            "abuse_score": d.get("abuseConfidenceScore"),
            "total_reports": d.get("totalReports"),
            "distinct_users": d.get("numDistinctUsers"),
            "last_reported": d.get("lastReportedAt"),
            "country": d.get("countryCode"),
            "isp": d.get("isp"),
            "usage_type": d.get("usageType"),
        }
    except Exception as e:
        result["error"] = str(e)
    return result
